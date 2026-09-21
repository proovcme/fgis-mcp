"""Loss-aware normalization of public API records. No edition or norm selection."""

import hashlib
import html
import json
import math
import re

from .network import SourceError


def clean(value):
    return " ".join(html.unescape(re.sub(r"<[^>]*>", " ", str(value or ""))).split())


def code_text(value):
    # Highlight tags can occur in the middle of a norm code.
    return re.sub(r"\s+", "", html.unescape(re.sub(r"<[^>]*>", "", str(value or ""))))


def number(value):
    if value is None:
        return None
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        result = float(text)
        return result if math.isfinite(result) else None
    except ValueError:
        return None


def array(record, key):
    value = record.get(key)
    if value is None or value == "":
        return []
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError as exc:
            raise SourceError(
                "SCHEMA_CHANGED", f"Invalid nested JSON in {key}; original response retained"
            ) from exc
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise SourceError("SCHEMA_CHANGED", f"Expected an array of objects in {key}")
    return value


def parse_hierarchy(raw_doc):
    """Parse hierarchical levels from source document name (collection, department, section, subsection, table)."""
    if not raw_doc:
        return {
            "collection": None,
            "department": None,
            "section": None,
            "subsection": None,
            "table": None,
            "full_path": [],
        }
    raw_str = str(raw_doc)
    if re.search(r"<br\s*/?>|\r?\n", raw_str, re.IGNORECASE):
        parts = [clean(p) for p in re.split(r"<br\s*/?>|\r?\n", raw_str, flags=re.IGNORECASE) if clean(p)]
    else:
        pattern = r"(?=(?:^|\s)(?:Сборник\s+\d+|Отдел\s+\d+|Раздел\s+\d+|Подраздел\s+\d+|Таблица\s+(?:(?:ГЭСН|ФЕР|ТЕР)\S*\s+)?\d+))"
        chunks = [clean(p) for p in re.split(pattern, raw_str, flags=re.IGNORECASE) if clean(p)]
        if len(chunks) <= 1:
            chunks = [
                clean(p)
                for p in re.split(
                    r"(?=(?:^|\s)(?:Сборник|Отдел|Раздел|Подраздел|Таблица)\s+)", raw_str, flags=re.IGNORECASE
                )
                if clean(p)
            ]
        parts = chunks if len(chunks) > 1 else ([clean(raw_str)] if clean(raw_str) else [])

    collection, department, section, subsection, table = None, None, None, None, None
    for p in parts:
        pl = p.lower()
        if pl.startswith("сборник") and not collection:
            collection = p
        elif pl.startswith("отдел") and not department:
            department = p
        elif pl.startswith("раздел") and not section:
            section = p
        elif pl.startswith("подраздел") and not subsection:
            subsection = p
        elif pl.startswith("таблица") and not table:
            table = p

    return {
        "collection": collection,
        "department": department,
        "section": section,
        "subsection": subsection,
        "table": table,
        "full_path": parts,
    }


def norm_cards(records, *, default_source: str = "online_api"):
    """Preserve every publication and every resource value, including non-numeric values."""
    cards = []
    for record in records:
        columns = array(record, "normTableJson")
        if not columns:
            raise SourceError("SCHEMA_CHANGED", "A norm search record has no norm columns")
        values = array(record, "normTableValueTableJson")
        works = array(record, "normCatalogWorkTableJson")
        digest = hashlib.sha256(json.dumps(record, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        raw_doc = record.get("documentName") or record.get("name")
        clean_doc = clean(raw_doc)
        hierarchy = parse_hierarchy(raw_doc)
        doc_guid = record.get("normLegalDocPublishedGuid")
        rec_id = record.get("id")
        raw_edition = record.get("edition")
        edition_val = clean(raw_edition) if raw_edition else None
        source_url = (
            f"https://fgiscs.minstroyrf.ru/api/NormLegalDocFilePublished/GetByGuid/{doc_guid}"
            if (doc_guid and default_source == "online_api")
            else None
        )
        source = {
            "record_id": rec_id,
            "document": clean_doc,
            "document_guid": doc_guid,
            "record_sha256": digest,
        }
        evidence = {
            "source": default_source,
            "source_type": "SearchEstimatedRates" if default_source == "online_api" else default_source,
            "document": clean_doc,
            "document_guid": doc_guid,
            "record_id": rec_id,
            "sha256": digest,
        }
        provenance = {
            "source": default_source,
            "document_guid": doc_guid,
            "source_url": source_url,
            "edition": edition_val,
            "sha256": digest,
            "record_id": rec_id,
        }
        family = clean(record.get("documentTypeName"))
        if not family and raw_doc:
            raw_upper = str(raw_doc).upper()
            if "ГЭСНМ" in raw_upper:
                family = "ГЭСНм"
            elif "ГЭСНР" in raw_upper:
                family = "ГЭСНр"
            elif "ГЭСНП" in raw_upper:
                family = "ГЭСНп"
            elif "ГЭСН" in raw_upper:
                family = "ГЭСН"
            elif "ФЕРМ" in raw_upper:
                family = "ФЕРм"
            elif "ФЕР" in raw_upper:
                family = "ФЕР"
        by_code = {}
        for col in columns:
            code = code_text(col.get("number") or col.get("Number"))
            if not code:
                raise SourceError("SCHEMA_CHANGED", "Norm column has no code")
            by_code[code] = {
                "norm_id": f"{digest}:{code}",
                "code": code,
                "family": family,
                "name": clean(col.get("name") or col.get("Name")),
                "unit": clean(col.get("meterName") or col.get("MeterName")),
                "hierarchy": hierarchy,
                "source": source,
                "evidence": evidence,
                "provenance": provenance,
                "document_guid": doc_guid,
                "record_id": rec_id,
                "edition": edition_val,
                "work_steps": [],
                "resources": [],
                "massa": None,
                "special_indicators": [],
                "warnings": [],
            }
        # Search may return one column but resource values for sibling norms. Retain those
        # siblings and their explicit names; do not invent missing units or merge editions.
        tree = {v.get("NormTablePartId"): v for v in values}
        for row in values:
            quantities = row.get("NormTablePartNormValueList") or []
            if not isinstance(quantities, list):
                raise SourceError("SCHEMA_CHANGED", "Resource quantities are not a list")
            ancestors, seen = [], set()
            parent = row.get("NormTablePartParentId")
            while parent is not None and parent in tree and parent not in seen:
                seen.add(parent)
                ancestors.insert(0, clean(tree[parent].get("Name")))
                parent = tree[parent].get("NormTablePartParentId")
            for quantity in quantities:
                code = code_text(quantity.get("NormNumber"))
                if not code:
                    raise SourceError("SCHEMA_CHANGED", "Resource quantity has no norm reference")
                if code not in by_code:
                    by_code[code] = {
                        "norm_id": f"{digest}:{code}",
                        "code": code,
                        "family": family,
                        "name": clean(quantity.get("NormName")),
                        "unit": None,
                        "hierarchy": hierarchy,
                        "source": source,
                        "evidence": evidence,
                        "provenance": provenance,
                        "document_guid": doc_guid,
                        "record_id": rec_id,
                        "edition": edition_val,
                        "work_steps": [],
                        "resources": [],
                        "massa": None,
                        "special_indicators": [],
                        "warnings": ["Norm column absent in search response; unit unavailable"],
                    }
                by_code[code]["resources"].append(
                    {
                        "part_id": row.get("NormTablePartId"),
                        "parent_part_id": row.get("NormTablePartParentId"),
                        "code": code_text(row.get("Cipher")),
                        "name": clean(row.get("Name")),
                        "unit": clean(row.get("UnitName")),
                        "category_path": ancestors,
                        "quantity": number(quantity.get("Value")),
                        "quantity_raw": quantity.get("Value"),
                    }
                )
        for work in works:
            code = code_text(work.get("NormNumber"))
            if code in by_code:
                by_code[code]["work_steps"].append(clean(work.get("Name")))

        for card in by_code.values():
            spec_indicators = []
            massa_indicator = None
            for res in card["resources"]:
                c_code = res.get("code") or ""
                c_name = res.get("name") or ""
                is_mass = c_code == "5-1" or "масса" in c_name.casefold()
                is_special = is_mass or c_code.startswith("5-")
                if is_special:
                    item = {
                        "code": c_code,
                        "name": c_name,
                        "unit": res.get("unit"),
                        "value": res.get("quantity"),
                    }
                    spec_indicators.append(item)
                    if is_mass and massa_indicator is None:
                        massa_indicator = item
            card["massa"] = massa_indicator
            card["special_indicators"] = spec_indicators

        cards.extend(by_code.values())
    return cards


PRICE_COLUMNS = [
    ("code", ("код ресурс",)),
    ("name", ("наименование строительного ресурс", "наименование ресурс")),
    ("unit", ("единица измер", "ед. измер", "ед.изм")),
    ("price_current", ("текущем уровне",)),
    ("price_release", ("отпускная цена",)),
    ("price_base", ("сметная цена",)),
    ("group_no", ("номер группы",)),
    ("group_name", ("наименование группы",)),
    ("index", ("индекс",)),
]


def price_rows(path):
    """Stream rows from all recognized worksheets; keep original values and sheet/row provenance."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    recognized = 0
    try:
        for sheet in book:
            mapping = None
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if mapping is None:
                    if row_no > 100:
                        break
                    headers = [clean(v).casefold() for v in row]
                    if not any("код ресурс" in h for h in headers):
                        continue
                    mapping = {}
                    for i, header in enumerate(headers):
                        for field, needles in PRICE_COLUMNS:
                            if field not in mapping.values() and any(n in header for n in needles):
                                mapping[i] = field
                                break
                    if not {"code", "name", "unit", "price_base"}.issubset(mapping.values()):
                        raise SourceError("SCHEMA_CHANGED", "Split-form header is missing required columns")
                    recognized += 1
                    continue
                raw = {field: row[i] if i < len(row) else None for i, field in mapping.items()}
                code = code_text(raw.get("code"))
                if not re.search(r"\d[.\d]*[-.]\d", code):
                    continue
                normalized = dict(raw)
                normalized.update(
                    code=code,
                    name=clean(raw.get("name")),
                    unit=clean(raw.get("unit")),
                    sheet=sheet.title,
                    row=row_no,
                )
                for field in ("price_base", "price_release", "price_current", "index"):
                    normalized[field] = number(raw.get(field))
                # Do not calculate or substitute published prices in an extraction service.
                normalized["raw_values"] = {k: str(v) if v is not None else None for k, v in raw.items()}
                yield normalized
        if not recognized:
            raise SourceError("SCHEMA_CHANGED", "No split-form worksheet recognized")
    finally:
        book.close()
